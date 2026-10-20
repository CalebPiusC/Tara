"""Read a CV from a file. Supported: PDF (pypdf), DOCX (python-docx), TXT/MD. Bounded and read-only.

PDFs without a text layer (scans) return no text and a warning; OCR is not included.
"""
from __future__ import annotations

import io
from pathlib import Path

MAX_BYTES = 10_000_000
MAX_PAGES = 30
MAX_CHARS = 60_000
SUFFIXES = (".pdf", ".docx", ".txt", ".md")


class CvError(ValueError):
    pass


def resolve_inside(base: Path, name: str) -> Path:
    """A file the owner dropped into the inbox. The model can name a file but never escape the directory."""
    if not name or "\x00" in name:
        raise CvError("no file name given")
    candidate = (base / name).resolve()
    root = base.resolve()
    if root != candidate and root not in candidate.parents:
        raise CvError("the file must be inside the candidate inbox")
    if not candidate.is_file():
        raise CvError("file not found in the candidate inbox")
    return candidate


def read_cv(path: Path) -> tuple[str, list[str]]:
    """Return (text, warnings). Raises CvError for unsupported or oversized files."""
    suffix = path.suffix.lower()
    if suffix not in SUFFIXES:
        raise CvError(f"unsupported file type {suffix or '(none)'}; use one of {list(SUFFIXES)}")
    size = path.stat().st_size
    if size > MAX_BYTES:
        raise CvError("file is larger than 10 MB")
    data = path.read_bytes()
    warnings: list[str] = []
    if suffix in (".txt", ".md"):
        text = data.decode("utf-8", errors="replace")
    elif suffix == ".docx":
        text = _docx(data)
    else:
        text, pdf_warnings = _pdf(data)
        warnings += pdf_warnings
    text = "\n".join(line.rstrip() for line in text.splitlines()).strip()
    if len(text) > MAX_CHARS:
        text = text[:MAX_CHARS]
        warnings.append(f"text was truncated to {MAX_CHARS} characters")
    if not text:
        warnings.append("no text could be extracted (scanned PDF or empty file?)")
    return text, warnings


def _docx(data: bytes) -> str:
    from docx import Document
    try:
        doc = Document(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001 - a corrupt file is the owner's problem to see, with a clear message
        raise CvError("could not open the DOCX file") from exc
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(parts)


def _pdf(data: bytes) -> tuple[str, list[str]]:
    from pypdf import PdfReader
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise CvError("the PDF is encrypted")
        pages = reader.pages
        warnings: list[str] = []
        if len(pages) > MAX_PAGES:
            warnings.append(f"only the first {MAX_PAGES} pages were read")
        text = "\n".join((pages[i].extract_text() or "") for i in range(min(len(pages), MAX_PAGES)))
        return text, warnings
    except CvError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise CvError("could not read the PDF file") from exc
